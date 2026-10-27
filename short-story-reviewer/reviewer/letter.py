"""Render the editorial letter from the letter stage's JSON.

Also checks the letter for quotations that are not in the story: in
diagnosis mode the letter should contain no replacement prose, and an
unsourced quotation is the usual sign of some.
"""
from __future__ import annotations

import html
import io
import re
from typing import Dict, List, Optional

import docx

KIND_LABELS = {
    "headline": "Headline",
    "pattern": "Pattern",
    "taste": "Taste",
    "query": "Query",
    "strength": "Strength",
}

_QUOTED = re.compile(r"[“\"]([^“”\"]{12,}?)[”\"]")


def _canon(text: str) -> str:
    text = re.sub(r"[‘’‛']", "'", text)
    text = re.sub(r"[“”\"]", '"', text)
    text = re.sub(r"—|–|--", "-", text)
    text = text.replace("…", "...")
    return re.sub(r"\s+", " ", text).strip().lower()


def unsourced_quotes(letter: Dict, story_text: str) -> List[str]:
    """Quoted passages of four words or more in the letter that the story does not contain."""
    story = _canon(story_text)
    found = []
    paragraphs = [p for s in letter.get("sections", []) for p in s.get("paragraphs", [])]
    for para in paragraphs:
        for m in _QUOTED.finditer(para):
            quote = m.group(1).strip().rstrip(",.;:!?")
            if len(quote.split()) < 4:
                continue
            pieces = [p for p in re.split(r"\s*(?:\.\.\.|…)\s*", quote) if p]
            if not all(_canon(p) in story for p in pieces):
                found.append(quote)
    return found


def to_markdown(meta: Dict, letter: Dict, located: List[Dict]) -> str:
    lines = [f"# Editorial letter: {meta['story_title']}", ""]
    lines.append(f"_{meta['subtitle']}_")
    lines.append("")
    if letter.get("salutation"):
        lines += [letter["salutation"], ""]
    for section in letter.get("sections", []):
        lines += [f"## {section['heading']}", ""]
        for para in section.get("paragraphs", []):
            lines += [para, ""]
    if letter.get("sign_off"):
        lines += [letter["sign_off"], ""]

    suggestions = letter.get("suggestions") or []
    if suggestions:
        lines += ["---", "", "## Suggested rewrites", "",
                  "These are suggestions only. They are not in the manuscript file; use, adapt or ignore them.", ""]
        for i, s in enumerate(suggestions, 1):
            lines.append(f"{i}. **Suggestion** for “{s['anchor']}”: {s['suggestion']}")
            if s.get("reason"):
                lines.append(f"   _Why:_ {s['reason']}")
        lines.append("")

    if located:
        lines += ["---", "", "## Notes located in the manuscript", ""]
        for n in located:
            label = KIND_LABELS.get(n.get("kind", ""), "Note")
            flag = "" if n.get("placed", True) else " _(could not be placed in the manuscript)_"
            lines.append(f"- “{n['anchor']}” — **{label}:** {n['note']}{flag}")
        lines.append("")
    return "\n".join(lines).rstrip() + "\n"


def to_html(meta: Dict, letter: Dict) -> str:
    e = html.escape
    out = [f"<h2>Editorial letter: {e(meta['story_title'])}</h2>",
           f"<p class=\"letter-meta\">{e(meta['subtitle'])}</p>"]
    if letter.get("salutation"):
        out.append(f"<p>{e(letter['salutation'])}</p>")
    for section in letter.get("sections", []):
        out.append(f"<h3>{e(section['heading'])}</h3>")
        out += [f"<p>{e(p)}</p>" for p in section.get("paragraphs", [])]
    if letter.get("sign_off"):
        out.append(f"<p>{e(letter['sign_off'])}</p>")
    suggestions = letter.get("suggestions") or []
    if suggestions:
        out.append("<h3>Suggested rewrites</h3>")
        out.append("<p class=\"letter-meta\">These are suggestions only. They are not in the manuscript file.</p>")
        out.append("<ol class=\"suggestions\">")
        for s in suggestions:
            reason = f"<br><em>Why:</em> {e(s['reason'])}" if s.get("reason") else ""
            out.append(f"<li><strong>Suggestion</strong> for “{e(s['anchor'])}”: "
                       f"{e(s['suggestion'])}{reason}</li>")
        out.append("</ol>")
    return "\n".join(out)


def to_docx(meta: Dict, letter: Dict, located: List[Dict]) -> bytes:
    document = docx.Document()
    document.add_heading(f"Editorial letter: {meta['story_title']}", level=0)
    sub = document.add_paragraph()
    sub.add_run(meta["subtitle"]).italic = True
    if letter.get("salutation"):
        document.add_paragraph(letter["salutation"])
    for section in letter.get("sections", []):
        document.add_heading(section["heading"], level=1)
        for para in section.get("paragraphs", []):
            document.add_paragraph(para)
    if letter.get("sign_off"):
        document.add_paragraph(letter["sign_off"])

    suggestions = letter.get("suggestions") or []
    if suggestions:
        document.add_heading("Suggested rewrites", level=1)
        document.add_paragraph(
            "These are suggestions only. They are not in the manuscript file; use, adapt or ignore them.")
        for s in suggestions:
            p = document.add_paragraph(style="List Number")
            p.add_run("Suggestion").bold = True
            p.add_run(f" for “{s['anchor']}”: {s['suggestion']}")
            if s.get("reason"):
                p.add_run(f" Why: {s['reason']}").italic = True

    if located:
        document.add_heading("Notes located in the manuscript", level=1)
        for n in located:
            p = document.add_paragraph(style="List Bullet")
            p.add_run(f"“{n['anchor']}” — ")
            p.add_run(f"{KIND_LABELS.get(n.get('kind', ''), 'Note')}: ").bold = True
            p.add_run(n["note"])
            if not n.get("placed", True):
                p.add_run(" (could not be placed in the manuscript)").italic = True

    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def subtitle(story_title: Optional[str], word_count: int, band: str, brief: Dict, date: str) -> str:
    mode = "diagnosis plus suggested rewrites" if brief.get("mode") == "rewrites" else "diagnosis only"
    return (f"{word_count:,} words ({band}) · {brief.get('edit_level', 'developmental')} edit · "
            f"{brief.get('draft_stage', 'early draft')} · {brief.get('conventions', 'UK')} conventions · "
            f"{mode} · {date}")
