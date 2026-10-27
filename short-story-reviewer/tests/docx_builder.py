"""Build .docx test manuscripts that look like Word made them.

Word rarely stores a paragraph as one run. Spell-check marks, revision IDs
and layout hints split the text, often mid-word. `fragment=True` imitates
that: it cuts runs at arbitrary points (keeping the same formatting), and
drops in w:proofErr and w:lastRenderedPageBreak markers, so the annotator's
merge step and single-run matching are exercised on realistic input.
"""
from __future__ import annotations

import io
import random
import re
from pathlib import Path
from typing import List, Tuple

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

STORIES = Path(__file__).resolve().parent / "fixtures" / "stories"


def load_fixture(slug: str) -> Tuple[str, List[str]]:
    """(title, paragraphs with _italic_ markers) for a fixture story."""
    raw = (STORIES / f"{slug}.md").read_text(encoding="utf-8")
    title = re.search(r"<!-- (.+?), by ", raw).group(1)
    body = re.sub(r"<!--.*?-->", "", raw, flags=re.S).strip()
    return title, [p.strip() for p in body.split("\n\n") if p.strip()]


def _segments(para: str) -> List[Tuple[str, bool]]:
    """Split a paragraph into (text, italic) pieces on _underscore_ markers."""
    out = []
    pos = 0
    for m in re.finditer(r"(?<![\w_])_(?=\S)(.+?)(?<=\S)_(?![\w_])", para):
        if m.start() > pos:
            out.append((para[pos:m.start()], False))
        out.append((m.group(1), True))
        pos = m.end()
    if pos < len(para):
        out.append((para[pos:], False))
    return out


def build_docx(paragraphs: List[str], title: str = "", fragment: bool = False, seed: int = 7) -> bytes:
    rng = random.Random(seed)
    document = docx.Document()
    if title:
        document.add_paragraph(title, style="Title")
    for para in paragraphs:
        p = document.add_paragraph()
        for text, italic in _segments(para):
            pieces = [text]
            if fragment and len(text) > 12:
                cuts = sorted(rng.sample(range(1, len(text)), k=min(3, len(text) - 1)))
                pieces = [text[a:b] for a, b in zip([0] + cuts, cuts + [len(text)])]
            for piece in pieces:
                run = p.add_run(piece)
                run.italic = italic or None
                if fragment and rng.random() < 0.3:
                    marker = OxmlElement("w:proofErr")
                    marker.set(qn("w:type"), "spellStart")
                    run._r.addprevious(marker)
                if fragment and rng.random() < 0.1:
                    run._r.insert(1, OxmlElement("w:lastRenderedPageBreak"))
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def fixture_docx(slug: str, fragment: bool = True) -> bytes:
    title, paras = load_fixture(slug)
    return build_docx(paras, title=title, fragment=fragment)
